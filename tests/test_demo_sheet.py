"""The shared spreadsheet renderer in demo/lib/sheet.py.

This file is byte-identical in all four portfolio repos, the same way
test_demo_fetch.py is, because the thing it tests is byte-identical in all
four. It imports by path rather than by package so it does not care which
piece it is running inside.

What is worth testing here is not that the HTML looks nice. It is the set of
promises the clip makes to a client:

  * the AFTER frame reads a real file, and fails loudly when there is none;
  * what it shows is what the file holds, including the number format;
  * a filtered view says it is filtered, by keeping the source file's own
    column letters and row numbers rather than renumbering to hide the gaps;
  * the command on screen is the command that produced the output under it.

`tools/demo_lib_drift.py` covers demo/, so it is what keeps the four copies of
sheet.py identical. It deliberately does not read the repo root, so it does not
see this file — the same gap test_demo_fetch.py sits in. Copy both together.
"""

from __future__ import annotations

import csv
import importlib.util
import subprocess
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parent.parent
SHEET = REPO / "demo" / "lib" / "sheet.py"


def _load():
    spec = importlib.util.spec_from_file_location("demo_sheet", SHEET)
    module = importlib.util.module_from_spec(spec)
    # Registered before it executes: @dataclass looks its own module up in
    # sys.modules while the class body is still being built, and a module that
    # is not there yet makes that lookup fail with an unrelated AttributeError.
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


sheet = _load()

try:  # openpyxl is optional in some pieces; the .xlsx tests skip without it
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

needs_xlsx = pytest.mark.skipif(openpyxl is None, reason="openpyxl not installed")

ROWS = [
    ["A-1", "Widget", "10.50", "yes"],
    ["A-2", "Gadget", "3.00", "no"],
    ["A-3", "Doohickey", "", "yes"],
    ["A-4", "Thing", "7.25", "no"],
]
HEADERS = ["sku", "name", "price", "flagged"]


@pytest.fixture
def csv_file(tmp_path: Path) -> Path:
    path = tmp_path / "out.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(ROWS)
    return path


@pytest.fixture
def xlsx_file(tmp_path: Path) -> Path:
    book = openpyxl.Workbook()
    book.remove(book.active)
    first = book.create_sheet("Clean")
    first.append(HEADERS)
    for row in ROWS:
        first.append(row[:2] + [float(row[2]) if row[2] else None, row[3]])
        first.cell(row=first.max_row, column=3).number_format = "0.00"
    book.create_sheet("Rejects").append(HEADERS)
    path = tmp_path / "book.xlsx"
    book.save(path)
    return path


class TestItMustReadARealFile:
    """The promise the whole clip rests on. A renderer that can be handed
    literal data can render a mock-up, which is the flaw being fixed."""

    def test_a_missing_file_raises_rather_than_rendering_nothing(self, tmp_path):
        with pytest.raises(FileNotFoundError) as caught:
            sheet.read_table(tmp_path / "never-written.csv")
        # The message has to say why it matters, not just ENOENT: whoever hits
        # this is mid-recording and needs to know the scene ran out of order.
        assert "nothing truthful to film" in str(caught.value)

    def test_an_unknown_extension_is_refused(self, tmp_path):
        path = tmp_path / "notes.txt"
        path.write_text("a,b\n1,2\n")
        with pytest.raises(ValueError, match="expected .csv or .xlsx"):
            sheet.read_table(path)

    def test_a_view_cannot_be_built_without_a_table(self):
        with pytest.raises(AttributeError):
            sheet.view({"headers": HEADERS, "rows": ROWS})


class TestReading:
    def test_csv_headers_and_rows(self, csv_file):
        table = sheet.read_table(csv_file)
        assert table.headers == HEADERS
        assert table.rows == ROWS
        assert table.sheet_names == ("out.csv",)
        assert table.unit == "rows"

    def test_a_short_row_is_padded_not_ragged(self, tmp_path):
        path = tmp_path / "ragged.csv"
        path.write_text("a,b,c\n1,2\n")
        table = sheet.read_table(path)
        assert table.rows == [["1", "2", ""]]

    def test_an_empty_file_is_an_error(self, tmp_path):
        path = tmp_path / "empty.csv"
        path.write_text("")
        with pytest.raises(ValueError, match="is empty"):
            sheet.read_table(path)

    def test_a_named_column_that_does_not_exist_says_what_does(self, csv_file):
        table = sheet.read_table(csv_file)
        with pytest.raises(KeyError) as caught:
            table.index_of("vendor")
        assert "sku" in str(caught.value)

    def test_base_makes_the_path_on_screen_relative(self, tmp_path):
        out = tmp_path / "out"
        out.mkdir()
        path = out / "clean.csv"
        path.write_text("a\n1\n")
        assert sheet.read_table(path, base=tmp_path).where == "out/"
        # A file at the repo root names the repo rather than printing "./".
        root = tmp_path / "invoices.csv"
        root.write_text("a\n1\n")
        assert sheet.read_table(root, base=tmp_path).where == f"{tmp_path.name}/"

    def test_no_base_does_not_leak_an_absolute_path(self, csv_file):
        where = sheet.read_table(csv_file).where
        assert not where.startswith("/")


@needs_xlsx
class TestReadingWorkbooks:
    def test_sheets_and_the_active_one(self, xlsx_file):
        table = sheet.read_table(xlsx_file)
        assert table.sheet_names == ("Clean", "Rejects")
        assert table.sheet == "Clean"
        assert sheet.read_table(xlsx_file, "Rejects").sheet == "Rejects"

    def test_a_sheet_that_does_not_exist_says_which_do(self, xlsx_file):
        with pytest.raises(KeyError) as caught:
            sheet.read_table(xlsx_file, "Changes")
        assert "Clean" in str(caught.value)

    def test_a_number_format_is_honoured(self, xlsx_file):
        """10.5 stored with a 0.00 format reads 10.50 in Excel, and 10.5 if you
        only look at the value. The CSV beside it says 10.50, so rendering the
        value alone would put a difference on screen that is not in the file."""
        table = sheet.read_table(xlsx_file)
        assert table.cell(0, "price") == "10.50"
        assert table.cell(1, "price") == "3.00"

    def test_an_empty_cell_is_empty_not_the_word_none(self, xlsx_file):
        assert sheet.read_table(xlsx_file).cell(2, "price") == ""


class TestNumberFormats:
    @pytest.mark.parametrize(
        "fmt,places",
        [("0.00", 2), ("0.000", 3), ("#,##0", 0), ("0", 0),
         ("General", None), ("", None), ('"$"#,##0.00', None),
         ("0.00%", None), ("[Red]0.00;0.00", None), ("yyyy-mm-dd", None)],
    )
    def test_only_plain_decimal_formats_are_followed(self, fmt, places):
        """A currency, percentage or date format is not guessed at: the cell
        falls back to its raw value rather than being dressed up wrongly."""
        assert sheet._decimal_places(fmt) == places

    def test_a_whole_float_loses_its_trailing_zero_when_unformatted(self):
        assert sheet._text(3.0) == "3"
        assert sheet._text(3.0, "0.00") == "3.00"

    def test_none_is_an_empty_cell(self):
        assert sheet._text(None) == ""


class TestChoosingWhatToShow:
    def test_column_letters_are_the_source_files(self):
        assert [sheet.column_letter(i) for i in (0, 1, 25, 26, 27, 51, 52)] == \
            ["A", "B", "Z", "AA", "AB", "AZ", "BA"]

    def test_picked_columns_keep_their_own_letters(self, csv_file):
        """Columns 0, 1 and 3 render as A, B and D — not A, B, C. Renumbering
        would hide that a column was skipped, which is the difference between
        a filtered view and a complete one."""
        table = sheet.read_table(csv_file)
        picked = sheet.view(table, ["sku", "name", "flagged"])
        html = sheet.grid_html(picked, step="AFTER", said="x")
        letters = [line for line in html.split("<th") if 'class="letter"' in line]
        assert ">A<" in html and ">B<" in html and ">D<" in html
        assert ">C<" not in "".join(letters)

    def test_rows_keep_their_real_sheet_row_number(self, csv_file):
        """Row 1 is the file's header, so the first data row is 2."""
        table = sheet.read_table(csv_file)
        html = sheet.grid_html(sheet.view(table, rows=[0, 3]), step="AFTER", said="x")
        assert '<td class="rownum">2</td>' in html
        assert '<td class="rownum">5</td>' in html
        assert '<td class="rownum">3</td>' not in html

    def test_a_gap_between_rows_is_marked(self, csv_file):
        table = sheet.read_table(csv_file)
        gapped = sheet.grid_html(sheet.view(table, rows=[0, 3]), step="AFTER", said="x")
        contiguous = sheet.grid_html(
            sheet.view(table, rows=[0, 1]), step="AFTER", said="x"
        )
        # On the row, not merely in the stylesheet — which always names it.
        assert '<tr class="skip' in gapped
        assert '<tr class="skip' not in contiguous

    def test_limit_takes_the_top_of_the_file(self, csv_file):
        table = sheet.read_table(csv_file)
        assert sheet.view(table, limit=2).rows == [0, 1]

    def test_limit_does_not_invent_rows_a_short_file_lacks(self, csv_file):
        assert sheet.view(sheet.read_table(csv_file), limit=99).rows == [0, 1, 2, 3]

    def test_rows_where_matches_a_value_or_a_predicate(self, csv_file):
        table = sheet.read_table(csv_file)
        assert sheet.rows_where(table, "flagged", "yes") == [0, 2]
        assert sheet.rows_where(table, "price", lambda v: not v) == [2]

    def test_head_and_keeps_file_order_and_caps(self, csv_file):
        table = sheet.read_table(csv_file)
        assert sheet.head_and(table, [3], head=2) == [0, 1, 3]
        assert sheet.head_and(table, [1], head=2) == [0, 1]       # no duplicate
        assert len(sheet.head_and(table, [3], head=2, limit=2)) == 2


class TestTinting:
    def test_a_tinted_row_carries_its_class(self, csv_file):
        table = sheet.read_table(csv_file)
        tinted = sheet.tint(sheet.view(table), [0, 2], "flag")
        html = sheet.grid_html(tinted, step="AFTER", said="x",
                               legend={"flag": "2 rows"})
        assert html.count("tint-flag") >= 2
        assert "2 rows" in html

    def test_an_unknown_tint_is_refused(self, csv_file):
        with pytest.raises(KeyError, match="no tint"):
            sheet.tint(sheet.view(sheet.read_table(csv_file)), [0], "purple")

    def test_tint_does_not_drop_the_column_widths(self, csv_file):
        """tint() rebuilds the View. It has to carry every field across, and
        `widths` is the one a positional rebuild silently loses."""
        table = sheet.read_table(csv_file)
        wide = sheet.view(table, widths={"name": 3.0})
        assert sheet.tint(wide, [0], "flag").widths == wide.widths

    def test_widths_reach_the_rendered_columns(self, csv_file):
        table = sheet.read_table(csv_file)
        plain = sheet.grid_html(sheet.view(table), step="AFTER", said="x")
        wide = sheet.grid_html(
            sheet.view(table, widths={"name": 3.0}), step="AFTER", said="x"
        )
        assert plain != wide
        assert "<col style=" in wide


class TestWhatTheFrameSays:
    def test_the_footer_counts_the_whole_file_not_the_view(self, csv_file):
        table = sheet.read_table(csv_file)
        html = sheet.grid_html(sheet.view(table, ["sku"], limit=1),
                               step="AFTER", said="x")
        assert "4 rows × 4 columns" in html
        assert "showing 1 of them, 1 of 4 columns" in html

    def test_a_complete_view_does_not_claim_to_be_filtered(self, csv_file):
        html = sheet.grid_html(sheet.view(sheet.read_table(csv_file)),
                               step="AFTER", said="x")
        assert "columns" in html
        assert "of 4 columns" not in html

    def test_the_real_filename_and_sheet_tabs_are_on_screen(self, csv_file):
        html = sheet.grid_html(sheet.view(sheet.read_table(csv_file)),
                               step="AFTER", said="x")
        assert "out.csv" in html

    def test_values_are_escaped(self, tmp_path):
        path = tmp_path / "x.csv"
        path.write_text('a\n"<script>alert(1)</script>"\n')
        html = sheet.grid_html(sheet.view(sheet.read_table(path)),
                               step="AFTER", said="x")
        assert "<script>" not in html
        assert "&lt;script&gt;" in html


class TestReadingADirectory:
    def test_it_lists_what_is_really_there(self, tmp_path):
        (tmp_path / "b.eml").write_text("second")
        (tmp_path / "a.eml").write_text("first")
        (tmp_path / "skip.txt").write_text("not matched")
        (tmp_path / "sub").mkdir()

        table = sheet.read_dir(tmp_path, "*.eml")
        assert [row[0] for row in table.rows] == ["a.eml", "b.eml"]
        assert table.headers == ["name", "type", "size", "modified"]
        assert table.rows[0][1] == "EML"
        assert table.rows[0][2] == "5 B"

    def test_a_listing_counts_files_not_rows(self, tmp_path):
        (tmp_path / "a.pdf").write_text("x")
        table = sheet.read_dir(tmp_path)
        assert table.unit == "files"
        html = sheet.grid_html(sheet.view(table), step="BEFORE", said="x")
        assert "1 files" in html

    def test_a_missing_directory_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            sheet.read_dir(tmp_path / "nope")

    def test_an_unknown_sort_is_refused(self, tmp_path):
        with pytest.raises(ValueError, match="no sort"):
            sheet.read_dir(tmp_path, sort="random")


class TestTheCommandBeat:
    def test_the_command_shown_is_the_command_run(self, tmp_path):
        command = sheet.run_command(
            [sys.executable, "-c", "print('42 rows written')"], cwd=tmp_path
        )
        assert command.stdout == "42 rows written"
        assert command.display.startswith("python -c")
        assert sys.executable not in command.display   # no machine paths on screen

    def test_stderr_is_kept_so_a_warning_is_not_hidden(self, tmp_path):
        command = sheet.run_command(
            [sys.executable, "-c",
             "import sys; sys.stderr.write('careful\\n'); print('done')"],
            cwd=tmp_path,
        )
        assert "careful" in command.stdout

    def test_a_failing_command_stops_the_recording(self, tmp_path):
        with pytest.raises(subprocess.CalledProcessError):
            sheet.run_command([sys.executable, "-c", "raise SystemExit(3)"],
                              cwd=tmp_path)

    def test_the_output_on_screen_is_the_captured_output(self, tmp_path):
        command = sheet.run_command(
            [sys.executable, "-c", "print('298 clean, 15 rejected')"], cwd=tmp_path
        )
        html = sheet.terminal_html(command, said="x")
        assert "298 clean, 15 rejected" in html

    def test_a_long_output_says_it_was_trimmed(self, tmp_path):
        command = sheet.run_command(
            [sys.executable, "-c", "print('\\n'.join(str(i) for i in range(30)))"],
            cwd=tmp_path,
        )
        html = sheet.terminal_html(command, said="x", max_lines=5)
        assert "25 more lines" in html
        assert ">5<" not in html.split('class="out"')[1][:200]


class TestNothingOnScreenNamesThisMachine:
    """THE-261: catalog-watch filmed `--out /home/<user>/rookery/projects/...`,
    which is our filesystem layout and a login name on the asset the Freelancer
    bid bodies link, and which ran off the right edge of the panel as well. The
    clip was otherwise fine — it encoded, it hit its byte count, it ended on the
    right frame — so nothing caught it but a human looking at a contact sheet.
    These are that human. They belong on `terminal_html` rather than on one
    scene because all four pieces film a command and any of them could grow the
    same defect tomorrow."""

    def _command(self, display: str):
        return sheet.Command(display=display, stdout="12 rows", returncode=0)

    def test_an_absolute_path_in_the_command_is_refused(self):
        with pytest.raises(ValueError, match="absolute path"):
            sheet.terminal_html(
                self._command("python watch.py --out /home/somebody/repo/out"),
                said="x",
            )

    def test_the_relative_form_of_the_same_command_is_fine(self):
        html = sheet.terminal_html(self._command("python watch.py --out out"),
                                   said="x")
        assert "--out out" in html

    def test_a_command_wider_than_the_panel_is_refused(self):
        too_long = "python clean.py " + "a" * sheet.MAX_COMMAND_CHARS
        with pytest.raises(ValueError, match="fit the panel"):
            sheet.terminal_html(self._command(too_long), said="x")

    def test_the_limit_is_the_width_that_actually_fits(self):
        """1280px frame, 26px of `main` padding and 26px of `.term` padding on
        each side, 15.5px DejaVu Sans Mono at 0.602em advance. The constant is
        a ceiling under that measured fit, not equal to it."""
        usable = sheet.VIEWPORT["width"] - 2 * 26 - 2 * 26
        fits = usable / (15.5 * 0.602)
        assert sheet.MAX_COMMAND_CHARS < fits

    def test_the_interpreter_path_is_not_what_trips_it(self, tmp_path):
        """run_command already rewrites argv[0] to `python`. If it stopped, the
        guard above would fire on every piece at once — so this pins the reason
        the guard stays quiet in normal use."""
        command = sheet.run_command([sys.executable, "-c", "print('ok')"],
                                    cwd=tmp_path)
        assert sheet.terminal_html(command, said="x")


class TestTheFourClipsShareOneShape:
    def test_the_holds_are_one_set_of_numbers(self):
        assert (sheet.HOLD_BEFORE, sheet.HOLD_COMMAND, sheet.HOLD_AFTER) == \
            (5.0, 3.6, 9.0)

    def test_the_holds_sit_inside_the_bands_the_company_rule_names(self):
        """COMPANY.md, Josue 2026-09-23: BEFORE 5-8s, AFTER 8-10s. THE-261 was
        allowed to take time out of the first third and explicitly not out of
        the AFTER frame, so this pins the floor that trimming may not cross."""
        assert 5.0 <= sheet.HOLD_BEFORE <= 8.0
        assert 8.0 <= sheet.HOLD_AFTER <= 10.0

    def test_the_whole_clip_fits_the_record_sh_budget(self):
        """record.sh rejects a clip over 35s. The shared beats have to leave
        room for a piece that splits one of them in two."""
        assert sheet.HOLD_BEFORE + sheet.HOLD_COMMAND + sheet.HOLD_AFTER < 35

    def test_the_typing_finishes_before_the_output_lands(self):
        assert sheet.OUTPUT_DELAY > sheet.TYPE_DELAY + sheet.TYPE_SECONDS
        assert sheet.OUTPUT_DELAY < sheet.HOLD_COMMAND

    def test_every_piece_films_at_the_same_size(self):
        assert sheet.VIEWPORT == {"width": 1280, "height": 720}
