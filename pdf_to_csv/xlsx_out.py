"""invoices.xlsx: the same table, in the file a client actually double-clicks.

A CSV is what a pipeline wants. A bookkeeper asked for the numbers, and what
they open is Excel — so the run writes both, from one `Table`, every time. The
CSV is not optional and is not diminished: `invoices.csv` is byte-for-byte
what it always was, and this file is written beside it.

Two things here are not decoration:

* **The numeric columns go back to being numbers.** A total stored as text is
  a total nobody can sum, which is most of why a client asked for the
  spreadsheet rather than the CSV. The number format keeps the two decimals on
  screen, so the cell still reads `1299.00` and does not quietly disagree with
  the CSV next to it.
* **The bytes are reproducible.** Two runs over the same PDFs produce an
  identical file, which is what lets `cmp` stand in for "trust me". See
  `_repack`.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

from .table import Table

SHEET_NAME = "Invoices"

#: Frozen so that two runs produce identical bytes. openpyxl would otherwise
#: stamp `docProps/core.xml` with the current time, and a workbook that differs
#: every run cannot be used as evidence that nothing changed.
EPOCH = datetime(1980, 1, 1, 0, 0, 0)

# Written as numbers rather than text. Anything not named here stays a string:
# `invoice_number` is the obvious one — "0042" is an identifier, and Excel eats
# the leading zero the moment it decides otherwise — and so is `invoice_date`,
# which is already ISO and sorts correctly as text in every locale.
NUMERIC_COLUMNS = frozenset(
    {"line_no", "quantity", "unit_price", "line_total", "subtotal", "tax", "total"}
)

# Money keeps its two decimals on screen. Without this the cell holds 1299 and
# Excel prints "1299", which disagrees with invoices.csv's "1299.00" over
# nothing: the value is identical, only the display was missing. `quantity` is
# deliberately not here — it is written `%g` in the CSV, so 2 is "2" and not
# "2.00", and General format spells it the same way.
MONEY_COLUMNS = frozenset({"unit_price", "line_total", "subtotal", "tax", "total"})
MONEY_FORMAT = "0.00"

# Kept close to the proportions the clip uses, so the sheet a client opens and
# the sheet in the recording are the same sheet. `description` and `issues` are
# prose and get clipped by the column rather than stretching it off the screen.
COLUMN_WIDTH = {"source_file": 32, "vendor": 24, "invoice_number": 16,
                "invoice_date": 13, "currency": 9, "line_no": 8,
                "description": 38, "quantity": 10, "unit_price": 12,
                "line_total": 12, "subtotal": 11, "tax": 10, "total": 11,
                "needs_review": 13, "issues": 46}
DEFAULT_WIDTH = 15


def _as_number(column: str, value: str):
    """A CSV cell as the number it spells, or unchanged when it is not one."""
    if column not in NUMERIC_COLUMNS or value in ("", None):
        return value
    try:
        text = str(value)
        return int(text) if text.lstrip("-").isdigit() else float(text)
    except ValueError:
        return value


def write_workbook(path: Path, table: Table) -> int:
    """The table, as one sheet. Returns the number of data rows written.

    Flagged rows — the documents where a field had to be left empty or did not
    add up — are tinted amber, which means the same thing here as the
    `needs_review` column means in the CSV: kept, but we would not vouch for
    it. They are on screen rather than filtered away, because a row a client
    cannot see is a row they will not check.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)

    book = Workbook()
    book.properties.created = EPOCH
    book.properties.modified = EPOCH
    book.properties.creator = "pdf-to-csv"
    book.properties.lastModifiedBy = "pdf-to-csv"

    sheet = book.active
    sheet.title = table.name or SHEET_NAME
    sheet.append(table.columns)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F3E4E")
    for cell in sheet[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")

    flag_fill = PatternFill("solid", fgColor="FFF4D6")
    money = [i for i, column in enumerate(table.columns) if column in MONEY_COLUMNS]
    for index, record in enumerate(table.records):
        sheet.append([_as_number(column, value)
                      for column, value in zip(table.columns, table.cells(record))])
        row = sheet[sheet.max_row]
        for position in money:
            if isinstance(row[position].value, (int, float)):
                row[position].number_format = MONEY_FORMAT
        if index in table.flagged:
            for cell in row:
                cell.fill = flag_fill

    for position, column in enumerate(table.columns, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = COLUMN_WIDTH.get(
            column, DEFAULT_WIDTH
        )
    # The header stays put when the client scrolls, and the filter arrows are
    # there because the first thing anyone does to a sheet like this is filter
    # it by needs_review.
    sheet.freeze_panes = "A2"
    if len(table):
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(table.columns))}{len(table) + 1}"
        )

    buffer = io.BytesIO()
    book.save(buffer)
    path.write_bytes(_repack(buffer.getvalue()))
    return len(table)


#: The two places an .xlsx records a wall clock. Measured on openpyxl 3.1.5:
#: `created` honours `book.properties`, and `modified` is refreshed to the save
#: time whatever the properties said — so pinning it has to happen here, after
#: the save, and the property above is not enough on its own.
_TIMESTAMP = re.compile(
    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)"
)
_EPOCH_XML = b"1980-01-01T00:00:00Z"


def _repack(data: bytes) -> bytes:
    """Rewrite an xlsx with every clock in it flattened.

    Two of them. An .xlsx is a zip, and `ZipFile.writestr` stamps each member
    with the current local time; it is also an Office document, and
    `docProps/core.xml` carries a modification timestamp that openpyxl
    refreshes on save. Either one makes two runs a second apart produce
    different bytes for identical content, which would put "run it twice and
    the output is unchanged" out of reach of `cmp` — the one check a client
    might actually perform. Order and contents are left exactly as openpyxl
    wrote them; only the clocks go.
    """
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            body = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                body = _TIMESTAMP.sub(rb"\g<1>" + _EPOCH_XML + rb"\g<2>", body)
            info = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = item.compress_type
            info.external_attr = item.external_attr
            info.create_system = 0
            target.writestr(info, body)
    return out.getvalue()
